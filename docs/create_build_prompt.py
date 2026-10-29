from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'SchoolNet_Build_Prompt.docx'
LOGO = ROOT / 'assets' / 'schoolnet_logo.jpg'
NAVY = '17345B'
TEAL = '11B6A5'
PALE = 'EAF6F5'
LIGHT = 'F3F6F9'
GRAY = '5D6B7C'


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=9.3):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    tr_pr.append(repeat)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement('w:cantSplit')
    tr_pr.append(no_split)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = True
    prevent_row_split(table.rows[0])
    for i, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], text, True, 'FFFFFF', 9)
        shade(table.rows[0].cells[i], NAVY)
    set_repeat_table_header(table.rows[0])
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, text in enumerate(row):
            set_cell_text(cells[i], str(text), False, None, 8.8)
            if ri % 2 == 1:
                shade(cells[i], LIGHT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def add_page_field(paragraph):
    r = paragraph.add_run('Page ')
    r.font.size = Pt(9)
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run = OxmlElement('w:r')
    run.extend([begin, instr, separate, text, end])
    paragraph._p.append(run)


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.68)
    sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
    styles = doc.styles
    normal = styles['Normal']; normal.font.name = 'Aptos'; normal.font.size = Pt(10.3); normal.font.color.rgb = RGBColor.from_string('263547')
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12
    for name, size, color in [('Title', 31, NAVY), ('Heading 1', 19, NAVY), ('Heading 2', 14, TEAL), ('Heading 3', 11, NAVY)]:
        s = styles[name]; s.font.name = 'Aptos Display'; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(13 if name != 'Title' else 0); s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = header.add_run('SCHOOLNET  /  IMPLEMENTATION PROMPT')
    rr.font.name = 'Aptos'; rr.font.size = Pt(8); rr.font.bold = True; rr.font.color.rgb = RGBColor.from_string(TEAL)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run('Source-aligned build instruction  •  ')
    rr.font.name='Aptos'; rr.font.size=Pt(9); rr.font.color.rgb=RGBColor.from_string(GRAY)
    add_page_field(footer)


def main():
    doc = Document()
    setup(doc)
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    if LOGO.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(14)
        p.add_run().add_picture(str(LOGO), width=Inches(1.45))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(10)
    r = p.add_run('SCHOOLNET'); r.bold = True; r.font.name = 'Aptos Display'; r.font.size = Pt(32); r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Application Build Prompt'); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run('Flutter Android app  ·  FastAPI REST backend  ·  Neon PostgreSQL'); r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_paragraph()
    add_table(doc, ['Prompt basis', 'Instruction'], [
        ['Master plan', 'SchoolNet_Comprehensive_PRD_Master_Blueprint.pdf is the product and architecture source of truth.'],
        ['Execution requirements', 'pasted_content.txt supplies the implementation detail and four-phase execution sequence.'],
        ['Brand asset', 'Use the user-supplied SchoolNet logo image as the app logo; preserve the image unchanged.'],
        ['Required stack', 'Flutter/Dart mobile client, Python FastAPI REST API, Neon PostgreSQL with tenant and role isolation.'],
    ])
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.35); p.paragraph_format.right_indent = Inches(.35)
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(10)
    p.add_run('MISSION  ').bold = True
    p.add_run('Connect school administration, teacher workflows, student achievement, private guidance support and community engagement—while keeping student data private across every screen and service.')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('CORE WORKSTREAMS  ·  Identity  ·  Portals  ·  Academic workflows  ·  Scoped communication  ·  Offline access')
    r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Prepared as a reusable, self-contained build prompt'); r.italic = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_page_break()

    doc.add_heading('Prompt contents', level=1)
    for item in [
        'Build instruction',
        '1. Product vision and platform',
        '2. Roles and identity lifecycle',
        '3. Global interface and experience',
        '4. Portal requirements — Administrator, Teacher, Student, G&C and People',
        '5. Messaging permissions and privacy',
        '6. Academic rating and grading rules',
        '7. Core domain model and API',
        '8. Build sequence and expected deliverables',
        '9. Source interpretation and open decisions',
        '10. Required quality bar',
    ]:
        bullet(doc, item)

    doc.add_heading('Build instruction', level=1)
    p = doc.add_paragraph()
    p.add_run('Act as the lead systems architect and full-stack engineering lead. Build SchoolNet as a modular, multi-tenant school-management and direct-communication application. ').bold = True
    p.add_run('Use this prompt together with the supplied PDF master plan. Treat the PDF as authoritative where details differ; preserve all compatible requirements from the pasted execution instructions. Do not silently replace Flutter/FastAPI/Neon with another stack.')

    doc.add_heading('1. Product vision and platform', level=1)
    doc.add_paragraph('Deliver an Android-first Flutter application for mobile and tablet, with a cross-platform-ready structure, backed by a Python FastAPI REST service and Neon PostgreSQL. SchoolNet serves day, boarding and hybrid schools. It combines administration, teaching workflows, student academic tracking, private guidance support and a public community gateway.')
    add_table(doc, ['Layer', 'Required implementation'], [
        ['Mobile client', 'Flutter SDK and Dart; use a clear architecture with models, views, controllers/state management and services. Provider or Riverpod is acceptable. Target Android phones and tablets first.'],
        ['API', 'Python FastAPI REST API. Implement JWT authentication, bcrypt password hashing, role-aware access controls, uploads, background work where needed, and exports.'],
        ['Database', 'Neon PostgreSQL with normalized relational tables and Row-Level Security policies isolating institutions, roles, classes and messages.'],
        ['Offline', 'Hive or SQLite cache for available announcements, marklists, directories and other approved read-only data. Show an offline indicator.'],
        ['Exports', 'Generate Excel `.xlsx` files with openpyxl and Word `.docx` reports with python-docx.'],
    ])

    doc.add_heading('2. Roles and identity lifecycle', level=1)
    add_table(doc, ['Role', 'Account lifecycle and core purpose'], [
        ['Administrator', 'Institutional account created by trusted provisioning; manages the school, users, permissions, announcements, messages, terms and archives.'],
        ['Teacher', 'Created by an administrator, individually or by roster import; manages assigned resources, classes and marks.'],
        ['Student', 'Created by an administrator, individually or by roster import; views own section and published academic results.'],
        ['Guidance & Counseling (G&C)', 'Optional institutional role/module; module is toggleable by the administrator and manages private student support threads.'],
        ['People / Community', 'Public self-registration with name and email; read-only public school information and announcements, with inquiry messages routed only to the administrator.'],
    ])
    bullet(doc, 'Institutional roles are created exclusively by administrators using a single-user form or CSV/Excel roster import. Generate a temporary ID/PIN and require a secure password reset at first login.')
    bullet(doc, 'Password recovery requests route to the administrator’s central inbox. Never reveal whether a submitted email belongs to an account.')
    bullet(doc, 'People accounts use open email-based self-registration. They do not receive general directory messaging rights.')

    doc.add_heading('3. Global interface and experience', level=1)
    bullet(doc, 'Use the supplied logo as the application logo; preserve it as the product brand asset. Apply a cohesive SchoolNet identity with accessible contrast and tablet/mobile responsive layouts.')
    bullet(doc, 'Provide a global search for public announcements, subject notes and permitted directory entries; role-specific bottom navigation; settings for password and local notification preferences; light/dark themes and configurable school accent colors.')
    bullet(doc, 'Show a subtle connectivity badge when offline and serve only data the signed-in user is authorized to access from local Hive/SQLite storage.')
    bullet(doc, 'Do not expose private student records or conversations through global search, push notifications, caches or public feeds.')

    doc.add_heading('4. Portal requirements', level=1)
    doc.add_heading('4.1 Administrator portal', level=2)
    for t in [
        'Roster management: create teachers and students individually or import CSV/Excel rosters; generate temporary credentials and force first-login reset.',
        'Class/section matrix: configure Grades 9–12 and divisions A, B, C; assign or swap Class Head Teachers and subject teachers.',
        'Teacher appraisals: administrator-only control for assigning/updating a 5-star teacher performance badge.',
        'Central messaging hub: separate inbox folders for Students, Teachers (including leave/allowance requests), Guidance & Counseling when enabled, and People.',
        'Global broadcast feed: administrator-only publishing of institution-wide announcements.',
        'About School CMS: edit history, contact information, school rules and Telegram, Instagram and Facebook links; support map coordinates as specified by the PDF.',
        'Seasonal panel: switch active term/season, toggle Guidance & Counseling, and set day/boarding mode.',
        'Annual roll-over: snapshot current marklists and ranks into a year-searchable Past Year Archive Vault, promote eligible cohorts (for example Grade 9 → Grade 10), and remap teachers/permissions. Preserve historical snapshots.',
    ]: bullet(doc, t)

    doc.add_heading('4.2 Teacher portal', level=2)
    for t in [
        'Show avatar, profile details, assigned subjects, administrator-assigned 5-star badge and inbox.',
        'Upload study notes, assignment PDFs and learning resources tagged by grade and subject.',
        'Show color-coded assigned class cards and identify sections where the teacher is Head Teacher.',
        'Provide a marklist grid for maximum exam points and per-student scores; validate marks and calculate totals, percentages, letter grades and section ranking.',
        'On saving/publishing results, recalculate the student’s 10-star academic badge from the exact formula in section 6.',
        'Provide announcements, teacher directory/shared notes, student directory with edit rights only for assigned classes, and a read-only About School hub.',
    ]: bullet(doc, t)

    doc.add_heading('4.3 Student portal', level=2)
    for t in [
        'Show avatar, name, assigned grade/section and public 10-star academic badge, calculated from published marks.',
        'Restrict section resources and marklists to the student’s enrolled grade/section.',
        'Provide a structured academic/personal difficulty form routed privately to the student’s Class Head Teacher.',
        'If G&C is enabled, provide a direct private support-messaging launcher on the student profile.',
        'Provide grade-level feeds visible only to that grade and section archives visible only to that exact section.',
        'Allow read-only browsing of public announcements, teacher directories and basic student lists; do not allow free-form DMs to general teachers or administrators.',
    ]: bullet(doc, t)

    doc.add_heading('4.4 Guidance & Counseling and People portals', level=2)
    bullet(doc, 'G&C is optional and administrator-toggleable. Put private student support threads within the G&C profile view; retain read-only access to global school feeds.')
    bullet(doc, 'People accounts self-register using name and email. Provide direct inquiry submission to the administrator’s People inbox subfolder and read-only access to public announcements and About School.')

    doc.add_heading('5. Messaging permissions and privacy', level=1)
    add_table(doc, ['Sender', 'Allowed destination', 'Restriction'], [
        ['Administrator', 'All roles, broadcasts and central inbox folders', 'Full school messaging authority, bounded to the administrator’s institution.'],
        ['Teacher', 'Assigned students and administrator', 'Teacher-initiated student DMs are limited to assigned students; no cross-school access.'],
        ['Student', 'G&C, if enabled; structured Head Teacher form', 'No free-form DM to general teachers or administrators.'],
        ['People / Community', 'Administrator’s People inbox folder', 'Inquiry must route to school administration only.'],
    ])
    bullet(doc, 'Enforce tenant, role, grade and section isolation in database queries and/or PostgreSQL RLS—not only by hiding UI controls.')
    bullet(doc, 'Scope every query to the authenticated institution. Authorize class operations against current enrollment, teacher assignment or Head Teacher status.')
    bullet(doc, 'Do not let students see another student’s marks, ranking details or private support threads unless explicitly authorized by the source requirements and school policy.')
    bullet(doc, 'Use least privilege for API/database credentials; hash passwords with bcrypt; protect temporary PINs and tokens; audit sensitive administrative changes.')

    doc.add_heading('6. Academic rating and grading rules', level=1)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.22); p.paragraph_format.right_indent = Inches(.22)
    r = p.add_run('Calculated Stars = (Σ Earned Student Marks ÷ Σ Maximum Subject Points) × 10'); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph('Examples from the master plan: 92% = 9.2 stars; 95% = 9.5 stars. Recompute and synchronize the badge when marks are saved/published. Do not confuse the student 10-star academic scale with the separate administrator-assigned teacher 5-star appraisal badge.')
    bullet(doc, 'Reject negative scores and scores greater than the maximum points. Do not divide by zero; return no rating when there are no eligible published points.')
    bullet(doc, 'Use the marklist maximum points in the denominator for each included subject. Keep the rating formula consistent across API, profile and exports.')

    doc.add_heading('7. Core domain model and API', level=1)
    doc.add_paragraph('Build a normalized schema around the PDF’s core concepts: institutions, users, classes_and_sections, marklists, student_performance, messages_and_posts and archive_vault. Add supporting relational tables when needed for enrollments, subjects, teaching assignments, individual scores, recovery requests, resources, attachments and audit events. Include institution_id on tenant-owned records.')
    bullet(doc, 'Use UUID identifiers, explicit relationships, constraints and indexes. Keep historical archive snapshots immutable and searchable by academic year.')
    bullet(doc, 'Expose REST endpoints for authentication, first-password reset, user creation/import, class/section setup, assignments, marks, feeds, messages, performance, exports and seasonal administration.')
    bullet(doc, 'Return only role-appropriate data. Protect bulk imports and file uploads with size/type validation, authorization and safe storage practices.')

    doc.add_heading('8. Build sequence and expected deliverables', level=1)
    numbered(doc, 'Database and backend: create migrations for the core schema; configure Neon PostgreSQL; implement bcrypt authentication, JWTs, role/tenant middleware, first-login reset and recovery routing; implement roster import, mark calculations, secure feed/messaging routes and Excel/Word export utilities.')
    numbered(doc, 'Flutter foundation: create /lib/models, /lib/views, /lib/controllers and /lib/services (or equivalent); implement role selection, login, first-time reset, public People signup, navigation, brand/logo and theme system.')
    numbered(doc, 'Portals: implement Administrator, Teacher, Student, optional G&C and People workflows. Keep functionality and API permissions aligned across client and server.')
    numbered(doc, 'Offline and release readiness: add Hive/SQLite cache and offline indicator; implement `.xlsx` and `.docx` export; test role/tenant privacy boundaries, accessibility, Android phone/tablet layouts, theme switching and migration/roll-over behavior.')
    bullet(doc, 'Deliver readable source code, setup instructions, environment-variable examples, SQL/Alembic migrations, API documentation, and tests. Do not claim a production deployment or Android binary unless it has actually been built and verified.')

    doc.add_heading('9. Source interpretation and open decisions', level=1)
    bullet(doc, 'The PDF permits Flask or FastAPI; use FastAPI for this build prompt. The PDF specifies Flutter, Neon PostgreSQL, Hive/SQLite, openpyxl and python-docx; retain those choices.')
    bullet(doc, 'The PDF describes six/several conceptual core tables, while the execution text lists a slightly different set; normalize the schema without dropping any required concept.')
    bullet(doc, 'The PDF and pasted text use both “Guidance & Counseling” and “Counselor”; use COUNSELOR as the persisted role code and display the full product name.')
    bullet(doc, 'The PDF’s sample formula and the pasted example agree: an aggregate 92% maps to 9.2 stars. Preserve the exact formula and examples.')
    bullet(doc, 'Define and document choices that the source does not settle—such as letter-grade thresholds, academic year advancement for Grade 12, data retention, attachment limits, and tie-breaking for section rank—rather than silently inventing school policy.')

    doc.add_heading('10. Required quality bar', level=1)
    for t in [
        'Maintain strict institution and role isolation at the API/database layer.',
        'Verify first-login password reset and prevent access to protected routes before reset completes.',
        'Verify students cannot send unauthorized free-form messages or retrieve another section’s feed/marks.',
        'Verify teacher edit rights are limited to assigned classes/subjects and admin operations remain institution-scoped.',
        'Verify Community inquiries route only to the administrator and public users see only public content.',
        'Verify offline cache never becomes an authorization bypass and is cleared appropriately on sign-out/account change.',
        'Verify Excel/Word outputs match the authorized scope and star calculation.',
        'Use the supplied logo as the app brand asset, not a substitute logo or generic icon.',
        'Verify cross-institution queries and exports are denied under RLS using a non-owner runtime role.',
        'Verify roster import rejects wrong file types, malformed rows, duplicate identities and oversized files safely.',
        'Verify migrations apply to an empty database and document a tested recovery/backup path before upgrades.',
        'Verify logs never contain passwords, temporary PINs, bearer tokens, private messages or unnecessary student data.',
        'Require HTTPS, managed secrets, backups, monitoring and retention controls before production launch.',
        'Test responsive layouts, accessible labels and navigation on Android phone and tablet form factors.',
        'Provide a safe local demo setup and trusted first-administrator bootstrap without any public admin-creation route.',
    ]: bullet(doc, t)

    doc.core_properties.title = 'SchoolNet Application Build Prompt'
    doc.core_properties.subject = 'Flutter and FastAPI implementation prompt based on the SchoolNet master blueprint'
    doc.core_properties.author = 'SchoolNet project'
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    main()
