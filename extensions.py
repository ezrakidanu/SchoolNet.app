from __future__ import annotations

from typing import Annotated

from fastapi import Depends, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy import select

from .main import (CurrentUser, Db, Enrollment, Marklist, Post, Role, Score, Section,
                   Scope, Subject, TeachingAssignment, User, app, io,
                   require_roles, require_section_access)

class SubjectCreate(BaseModel):
    name: str = Field(min_length=1, max_length=120)

class AssignmentCreate(BaseModel):
    teacher_id: str
    section_id: str
    subject_id: str

class TeacherRating(BaseModel):
    stars: float = Field(ge=0, le=5)

class StudentDirectMessage(BaseModel):
    content: str = Field(min_length=1, max_length=10000)
    title: str | None = Field(default=None, max_length=180)

@app.post("/admin/subjects", status_code=201)
def create_subject(payload: SubjectCreate, user: Annotated[User, Depends(require_roles(Role.ADMIN))], db: Db) -> dict:
    subject = Subject(institution_id=user.institution_id, name=payload.name.strip())
    db.add(subject); db.commit()
    return {"id": subject.id, "name": subject.name}

@app.post("/admin/assignments", status_code=201)
def create_assignment(payload: AssignmentCreate, user: Annotated[User, Depends(require_roles(Role.ADMIN))], db: Db) -> dict:
    teacher = db.scalar(select(User).where(User.id == payload.teacher_id, User.institution_id == user.institution_id, User.role == Role.TEACHER.value))
    section = db.scalar(select(Section).where(Section.id == payload.section_id, Section.institution_id == user.institution_id))
    subject = db.scalar(select(Subject).where(Subject.id == payload.subject_id, Subject.institution_id == user.institution_id))
    if not teacher or not section or not subject:
        raise HTTPException(status_code=404, detail="Teacher, section or subject not found in this school")
    exists = db.scalar(select(TeachingAssignment.id).where(TeachingAssignment.teacher_id == teacher.id, TeachingAssignment.section_id == section.id, TeachingAssignment.subject_id == subject.id))
    if exists:
        raise HTTPException(status_code=409, detail="Assignment already exists")
    item = TeachingAssignment(institution_id=user.institution_id, teacher_id=teacher.id, section_id=section.id, subject_id=subject.id)
    db.add(item); db.commit()
    return {"id": item.id, "teacher_id": item.teacher_id, "section_id": item.section_id, "subject_id": item.subject_id}

@app.put("/admin/teachers/{teacher_id}/rating")
def set_teacher_rating(teacher_id: str, payload: TeacherRating, user: Annotated[User, Depends(require_roles(Role.ADMIN))], db: Db) -> dict:
    teacher = db.scalar(select(User).where(User.id == teacher_id, User.institution_id == user.institution_id, User.role == Role.TEACHER.value))
    if teacher is None:
        raise HTTPException(status_code=404, detail="Teacher not found in this school")
    teacher.teacher_stars = payload.stars
    db.commit()
    return {"teacher_id": teacher.id, "teacher_stars": teacher.teacher_stars}

@app.get("/marklists/section/{section_id}")
def section_marklists(section_id: str, user: CurrentUser, db: Db) -> list[dict]:
    require_section_access(db, user, section_id)
    query = select(Marklist).where(Marklist.institution_id == user.institution_id, Marklist.section_id == section_id)
    if user.role == Role.STUDENT.value:
        query = query.where(Marklist.published.is_(True))
    if user.role == Role.TEACHER.value:
        assigned_subjects = select(TeachingAssignment.subject_id).where(TeachingAssignment.teacher_id == user.id, TeachingAssignment.section_id == section_id)
        query = query.where((Marklist.teacher_id == user.id) | Marklist.subject_id.in_(assigned_subjects))
    return [{"id": m.id, "title": m.title, "subject_id": m.subject_id, "max_points": m.max_points,
             "term": m.term, "academic_year": m.academic_year, "published": m.published}
            for m in db.scalars(query.order_by(Marklist.created_at.desc())).all()]

@app.post("/messages/student/{student_id}", status_code=201)
def send_student_message(student_id: str, payload: StudentDirectMessage, user: Annotated[User, Depends(require_roles(Role.ADMIN, Role.TEACHER))], db: Db) -> dict:
    student = db.scalar(select(User).where(User.id == student_id, User.institution_id == user.institution_id, User.role == Role.STUDENT.value))
    if student is None:
        raise HTTPException(status_code=404, detail="Student not found in this school")
    if user.role == Role.TEACHER.value:
        enrollment = db.scalar(select(Enrollment).where(Enrollment.student_id == student.id))
        if enrollment is None:
            raise HTTPException(status_code=404, detail="Student has no current section")
        assignment = db.scalar(select(TeachingAssignment.id).where(TeachingAssignment.teacher_id == user.id, TeachingAssignment.section_id == enrollment.section_id))
        section = db.get(Section, enrollment.section_id)
        if not assignment and (not section or section.head_teacher_id != user.id):
            raise HTTPException(status_code=403, detail="Teachers may only message assigned students")
    message = Post(institution_id=user.institution_id, sender_id=user.id, recipient_id=student.id, scope=Scope.DM.value,
                   category="STAFF_TO_STUDENT", title=payload.title, content=payload.content)
    db.add(message); db.commit()
    return {"message": "Private message sent to student."}

@app.get("/messages/inbox")
def personal_inbox(user: CurrentUser, db: Db) -> list[dict]:
    rows = db.scalars(select(Post).where(Post.institution_id == user.institution_id, Post.scope == Scope.DM.value,
                        Post.recipient_id == user.id).order_by(Post.created_at.desc()).limit(200)).all()
    return [{"id": m.id, "sender_id": m.sender_id, "category": m.category, "title": m.title,
             "content": m.content, "created_at": m.created_at.isoformat()} for m in rows]

@app.get("/exports/students/{student_id}.docx")
def export_student_report(student_id: str, user: CurrentUser, db: Db):
    student = db.scalar(select(User).where(User.id == student_id, User.institution_id == user.institution_id, User.role == Role.STUDENT.value))
    if student is None:
        raise HTTPException(status_code=404, detail="Student not found in this school")
    if user.role not in {Role.ADMIN.value, Role.TEACHER.value, Role.STUDENT.value}:
        raise HTTPException(status_code=403, detail="This role cannot export student reports")
    if user.role == Role.STUDENT.value and user.id != student.id:
        raise HTTPException(status_code=403, detail="Students may export only their own report")
    if user.role == Role.TEACHER.value:
        enrollment = db.scalar(select(Enrollment).where(Enrollment.student_id == student.id))
        if enrollment is None:
            raise HTTPException(status_code=404, detail="Student has no current section")
        require_section_access(db, user, enrollment.section_id)
    rows = db.execute(select(Marklist.title, Marklist.term, Marklist.max_points, Score.earned_points)
                      .join(Score, Score.marklist_id == Marklist.id)
                      .where(Score.student_id == student.id, Marklist.institution_id == user.institution_id,
                             Marklist.published.is_(True)).order_by(Marklist.academic_year, Marklist.term, Marklist.title)).all()
    earned, maximum = sum(float(r.earned_points) for r in rows), sum(float(r.max_points) for r in rows)
    stars = round(earned / maximum * 10, 2) if maximum else None
    from docx import Document as WordDocument
    from docx.shared import Inches, Pt, RGBColor
    report = WordDocument()
    report.sections[0].top_margin = Inches(.7); report.sections[0].bottom_margin = Inches(.7)
    title = report.add_heading('SchoolNet · Student Report Card', 0)
    title.runs[0].font.color.rgb = RGBColor(23, 52, 91)
    report.add_paragraph(f'Student: {student.full_name}')
    report.add_paragraph(f'Academic average: {round(earned / maximum * 100, 2) if maximum else "Not available"}%')
    report.add_paragraph(f'10-star academic rating: {stars if stars is not None else "Not available"}')
    table = report.add_table(rows=1, cols=4); table.style = 'Light Shading Accent 1'
    for cell, text_value in zip(table.rows[0].cells, ['Assessment', 'Term', 'Earned', 'Maximum']): cell.text = text_value
    for title_text, term_text, max_points, points in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, [title_text, term_text, f'{points:g}', f'{max_points:g}']): cell.text = str(value)
    output = io.BytesIO(); report.save(output); output.seek(0)
    from fastapi.responses import StreamingResponse
    filename = f"schoolnet-report-{student.id[:8]}.docx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})
