from __future__ import annotations

import csv
import io
import os
import secrets
import uuid
from datetime import datetime, timedelta, timezone
from enum import Enum
from typing import Annotated, Generator

import bcrypt
import jwt
from fastapi import Depends, FastAPI, File, HTTPException, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from openpyxl import Workbook
from pydantic import BaseModel, EmailStr, Field
from sqlalchemy import (Boolean, DateTime, Float, ForeignKey, Integer, String, Text,
                        UniqueConstraint, create_engine, func, select, text)
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, sessionmaker

DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./schoolnet_dev.db")
JWT_SECRET = os.getenv("JWT_SECRET", "local-development-only-change-me")
JWT_ALGORITHM = "HS256"
TOKEN_TTL_MINUTES = int(os.getenv("TOKEN_TTL_MINUTES", "720"))

if not DATABASE_URL.startswith("sqlite") and (
    JWT_SECRET == "local-development-only-change-me"
    or JWT_SECRET.lower().startswith(("replace-with-", "change-me"))
    or len(JWT_SECRET) < 32
):
    raise RuntimeError("Set a random JWT_SECRET with at least 32 characters for PostgreSQL deployments.")

connect_args = {"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}
engine = create_engine(DATABASE_URL, connect_args=connect_args, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)


class Base(DeclarativeBase):
    pass


class Role(str, Enum):
    ADMIN = "ADMIN"
    TEACHER = "TEACHER"
    STUDENT = "STUDENT"
    COUNSELOR = "COUNSELOR"
    COMMUNITY = "COMMUNITY"


class Scope(str, Enum):
    GLOBAL = "GLOBAL"
    GRADE = "GRADE"
    SECTION = "SECTION"
    DM = "DM"


class Institution(Base):
    __tablename__ = "institutions"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    name: Mapped[str] = mapped_column(String(180), nullable=False)
    slug: Mapped[str] = mapped_column(String(100), unique=True, nullable=False)
    has_guidance_counseling: Mapped[bool] = mapped_column(Boolean, default=False, nullable=False)
    is_boarding_school: Mapped[bool] = mapped_column(Boolean, default=False, nullable=False)
    logo_url: Mapped[str | None] = mapped_column(String(500))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class User(Base):
    __tablename__ = "users"
    __table_args__ = (UniqueConstraint("institution_id", "email", name="uq_user_school_email"),)
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    full_name: Mapped[str] = mapped_column(String(180), nullable=False)
    email: Mapped[str] = mapped_column(String(320), nullable=False, index=True)
    role: Mapped[str] = mapped_column(String(24), nullable=False)
    password_hash: Mapped[str] = mapped_column(String(255), nullable=False)
    force_password_change: Mapped[bool] = mapped_column(Boolean, default=True, nullable=False)
    teacher_stars: Mapped[float | None] = mapped_column(Float)
    avatar_url: Mapped[str | None] = mapped_column(String(500))
    active: Mapped[bool] = mapped_column(Boolean, default=True, nullable=False)


class Section(Base):
    __tablename__ = "classes_and_sections"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    grade_level: Mapped[int] = mapped_column(Integer, nullable=False)
    section_name: Mapped[str] = mapped_column(String(20), nullable=False)
    head_teacher_id: Mapped[str | None] = mapped_column(ForeignKey("users.id"))
    academic_year: Mapped[str] = mapped_column(String(16), nullable=False)
    __table_args__ = (UniqueConstraint("institution_id", "grade_level", "section_name", "academic_year", name="uq_school_section_year"),)


class Enrollment(Base):
    __tablename__ = "enrollments"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    student_id: Mapped[str] = mapped_column(ForeignKey("users.id"), nullable=False, index=True)
    section_id: Mapped[str] = mapped_column(ForeignKey("classes_and_sections.id"), nullable=False, index=True)
    academic_year: Mapped[str] = mapped_column(String(16), nullable=False)
    __table_args__ = (UniqueConstraint("student_id", "academic_year", name="uq_student_year_enrollment"),)


class Subject(Base):
    __tablename__ = "subjects"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    name: Mapped[str] = mapped_column(String(120), nullable=False)


class TeachingAssignment(Base):
    __tablename__ = "teaching_assignments"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    teacher_id: Mapped[str] = mapped_column(ForeignKey("users.id"), nullable=False)
    section_id: Mapped[str] = mapped_column(ForeignKey("classes_and_sections.id"), nullable=False)
    subject_id: Mapped[str] = mapped_column(ForeignKey("subjects.id"), nullable=False)


class Marklist(Base):
    __tablename__ = "marklists"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    section_id: Mapped[str] = mapped_column(ForeignKey("classes_and_sections.id"), nullable=False, index=True)
    subject_id: Mapped[str] = mapped_column(ForeignKey("subjects.id"), nullable=False)
    teacher_id: Mapped[str] = mapped_column(ForeignKey("users.id"), nullable=False)
    title: Mapped[str] = mapped_column(String(180), nullable=False)
    max_points: Mapped[float] = mapped_column(Float, nullable=False)
    term: Mapped[str] = mapped_column(String(60), nullable=False)
    academic_year: Mapped[str] = mapped_column(String(16), nullable=False)
    published: Mapped[bool] = mapped_column(Boolean, default=False, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class Score(Base):
    __tablename__ = "student_scores"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    marklist_id: Mapped[str] = mapped_column(ForeignKey("marklists.id", ondelete="CASCADE"), nullable=False, index=True)
    student_id: Mapped[str] = mapped_column(ForeignKey("users.id"), nullable=False, index=True)
    earned_points: Mapped[float] = mapped_column(Float, nullable=False)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), onupdate=lambda: datetime.now(timezone.utc))
    __table_args__ = (UniqueConstraint("marklist_id", "student_id", name="uq_score_student_marklist"),)


class Post(Base):
    __tablename__ = "messages_and_posts"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    sender_id: Mapped[str] = mapped_column(ForeignKey("users.id"), nullable=False)
    scope: Mapped[str] = mapped_column(String(16), nullable=False)
    grade_level: Mapped[int | None] = mapped_column(Integer)
    section_id: Mapped[str | None] = mapped_column(ForeignKey("classes_and_sections.id"))
    recipient_id: Mapped[str | None] = mapped_column(ForeignKey("users.id"))
    category: Mapped[str] = mapped_column(String(40), default="POST", nullable=False)
    title: Mapped[str | None] = mapped_column(String(180))
    content: Mapped[str] = mapped_column(Text, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class RecoveryRequest(Base):
    __tablename__ = "recovery_requests"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    user_email: Mapped[str] = mapped_column(String(320), nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class ArchiveSnapshot(Base):
    __tablename__ = "archive_vault"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    institution_id: Mapped[str] = mapped_column(ForeignKey("institutions.id"), nullable=False, index=True)
    academic_year: Mapped[str] = mapped_column(String(16), nullable=False)
    snapshot_json: Mapped[str] = mapped_column(Text, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class LoginInput(BaseModel):
    institution_slug: str
    email: EmailStr
    password: str
    role: Role


class CommunitySignup(BaseModel):
    institution_slug: str
    full_name: str = Field(min_length=2, max_length=180)
    email: EmailStr
    password: str = Field(min_length=10, max_length=128)


class PasswordChange(BaseModel):
    new_password: str = Field(min_length=10, max_length=128)


class UserCreate(BaseModel):
    full_name: str = Field(min_length=2, max_length=180)
    email: EmailStr
    role: Role
    section_id: str | None = None


class SectionCreate(BaseModel):
    grade_level: int = Field(ge=9, le=12)
    section_name: str = Field(min_length=1, max_length=20)
    academic_year: str = Field(min_length=4, max_length=16)
    head_teacher_id: str | None = None


class MarklistCreate(BaseModel):
    section_id: str
    subject_id: str
    title: str = Field(min_length=1, max_length=180)
    max_points: float = Field(gt=0)
    term: str
    academic_year: str


class ScoreInput(BaseModel):
    scores: dict[str, float]
    publish: bool = False


class PostCreate(BaseModel):
    scope: Scope
    title: str | None = Field(default=None, max_length=180)
    content: str = Field(min_length=1, max_length=10000)
    grade_level: int | None = None
    section_id: str | None = None


class MessageInput(BaseModel):
    content: str = Field(min_length=1, max_length=10000)
    category: str = "INQUIRY"


bearer = HTTPBearer(auto_error=False)
app = FastAPI(title="SchoolNet API", version="0.1.0", description="Multi-tenant school management API MVP")
app.add_middleware(CORSMiddleware, allow_origins=os.getenv("CORS_ORIGINS", "http://localhost:3000,http://localhost:8080").split(","), allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


@app.on_event("startup")
def on_startup() -> None:
    # Local SQLite convenience only; production schema changes are migrations.
    if DATABASE_URL.startswith("sqlite"):
        Base.metadata.create_all(bind=engine)


def get_db() -> Generator[Session, None, None]:
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


Db = Annotated[Session, Depends(get_db)]


def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt(rounds=12)).decode("utf-8")


def verify_password(password: str, stored: str) -> bool:
    try:
        return bcrypt.checkpw(password.encode("utf-8"), stored.encode("utf-8"))
    except ValueError:
        return False


def set_tenant_context(db: Session, tenant_id: str) -> None:
    if not DATABASE_URL.startswith("sqlite"):
        db.execute(text("SELECT set_config('app.current_institution', :tenant, true)"), {"tenant": tenant_id})


def issue_token(user: User) -> str:
    now = datetime.now(timezone.utc)
    return jwt.encode({"sub": user.id, "tenant": user.institution_id, "role": user.role,
                       "iat": now, "exp": now + timedelta(minutes=TOKEN_TTL_MINUTES)}, JWT_SECRET, algorithm=JWT_ALGORITHM)


def get_current_user(credentials: Annotated[HTTPAuthorizationCredentials | None, Depends(bearer)], db: Db) -> User:
    if credentials is None:
        raise HTTPException(status_code=401, detail="Bearer token required")
    try:
        claims = jwt.decode(credentials.credentials, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user_id, tenant_id = claims.get("sub"), claims.get("tenant")
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Invalid or expired token")
    if not user_id or not tenant_id:
        raise HTTPException(status_code=401, detail="Invalid token claims")
    # Set the tenant context consumed by the PostgreSQL RLS policies in migrations/001_core.sql.
    set_tenant_context(db, tenant_id)
    user = db.scalar(select(User).where(User.id == user_id, User.institution_id == tenant_id, User.active.is_(True)))
    if user is None:
        raise HTTPException(status_code=401, detail="Account is inactive or no longer exists")
    if user.force_password_change:
        raise HTTPException(status_code=403, detail="Change your temporary password before using SchoolNet")
    return user


CurrentUser = Annotated[User, Depends(get_current_user)]


def require_roles(*roles: Role):
    def guard(user: CurrentUser) -> User:
        if user.role not in {r.value for r in roles}:
            raise HTTPException(status_code=403, detail="Role is not permitted to perform this action")
        return user
    return guard


def user_summary(user: User) -> dict:
    return {"id": user.id, "institution_id": user.institution_id, "full_name": user.full_name,
            "email": user.email, "role": user.role, "force_password_change": user.force_password_change,
            "teacher_stars": user.teacher_stars}


def student_section(db: Session, user: User) -> Section | None:
    enrollment = db.scalar(select(Enrollment).where(Enrollment.student_id == user.id).order_by(Enrollment.academic_year.desc()))
    return db.get(Section, enrollment.section_id) if enrollment else None


def require_section_access(db: Session, user: User, section_id: str) -> None:
    if user.role == Role.ADMIN.value:
        return
    if user.role == Role.TEACHER.value:
        assignment = db.scalar(select(TeachingAssignment.id).where(TeachingAssignment.teacher_id == user.id, TeachingAssignment.section_id == section_id))
        section = db.get(Section, section_id)
        if (section and section.head_teacher_id == user.id) or assignment:
            return
    if user.role == Role.STUDENT.value:
        enrolled = db.scalar(select(Enrollment.id).where(Enrollment.student_id == user.id, Enrollment.section_id == section_id))
        if enrolled:
            return
    raise HTTPException(status_code=403, detail="You do not have access to this class or section")


@app.get("/health")
def health() -> dict:
    return {"status": "ok", "service": "schoolnet-api", "version": app.version}


@app.post("/auth/login")
def login(payload: LoginInput, db: Db) -> dict:
    institution = db.scalar(select(Institution).where(Institution.slug == payload.institution_slug))
    if institution is None:
        raise HTTPException(status_code=401, detail="Invalid institution or credentials")
    set_tenant_context(db, institution.id)
    user = db.scalar(select(User).where(User.institution_id == institution.id, User.email == payload.email.lower(), User.role == payload.role.value))
    if user is None or not user.active or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Invalid institution or credentials")
    return {"access_token": issue_token(user), "token_type": "bearer", "user": user_summary(user)}


@app.post("/auth/register-community", status_code=201)
def register_community(payload: CommunitySignup, db: Db) -> dict:
    institution = db.scalar(select(Institution).where(Institution.slug == payload.institution_slug))
    if institution is None:
        raise HTTPException(status_code=404, detail="School not found")
    set_tenant_context(db, institution.id)
    email = payload.email.lower()
    if db.scalar(select(User.id).where(User.institution_id == institution.id, User.email == email)):
        raise HTTPException(status_code=409, detail="An account with this email already exists for this school")
    user = User(institution_id=institution.id, full_name=payload.full_name.strip(), email=email,
                role=Role.COMMUNITY.value, password_hash=hash_password(payload.password), force_password_change=False)
    db.add(user)
    db.commit()
    return {"access_token": issue_token(user), "token_type": "bearer", "user": user_summary(user)}


@app.post("/auth/change-password")
def change_password(payload: PasswordChange, credentials: Annotated[HTTPAuthorizationCredentials | None, Depends(bearer)], db: Db) -> dict:
    if credentials is None:
        raise HTTPException(status_code=401, detail="Bearer token required")
    try:
        claims = jwt.decode(credentials.credentials, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        if claims.get("tenant"):
            set_tenant_context(db, claims["tenant"])
        user = db.scalar(select(User).where(User.id == claims.get("sub"), User.institution_id == claims.get("tenant")))
    except jwt.PyJWTError:
        user = None
    if user is None:
        raise HTTPException(status_code=401, detail="Invalid or expired token")
    user.password_hash = hash_password(payload.new_password)
    user.force_password_change = False
    db.commit()
    return {"message": "Password updated. Sign in again with your new password."}


@app.post("/auth/recovery-request", status_code=202)
def recovery_request(institution_slug: str, email: EmailStr, db: Db) -> dict:
    institution = db.scalar(select(Institution).where(Institution.slug == institution_slug))
    if institution is None:
        # Keep response generic to avoid disclosing institution/account membership.
        return {"message": "If the school exists, the request has been routed to its administrator."}
    set_tenant_context(db, institution.id)
    db.add(RecoveryRequest(institution_id=institution.id, user_email=email.lower()))
    db.commit()
    return {"message": "If the school exists, the request has been routed to its administrator."}


@app.get("/me")
def me(user: CurrentUser) -> dict:
    return user_summary(user)


@app.post("/admin/users", status_code=201)
def create_user(payload: UserCreate, user: Annotated[User, Depends(require_roles(Role.ADMIN))], db: Db) -> dict:
    if payload.role == Role.COMMUNITY:
        raise HTTPException(status_code=400, detail="Community users self-register")
    email = payload.email.lower()
    if db.scalar(select(User.id).where(User.institution_id == user.institution_id, User.email == email)):
        raise HTTPException(status_code=409, detail="Email already exists in this school")
    temp_pin = secrets.token_urlsafe(9)
    new_user = User(institution_id=user.institution_id, full_name=payload.full_name.strip(), email=email,
                    role=payload.role.value, password_hash=hash_password(temp_pin), force_password_change=True)
    db.add(new_user)
    db.flush()
    if payload.role == Role.STUDENT and payload.section_id:
        section = db.scalar(select(Section).where(Section.id == payload.section_id, Section.institution_id == user.institution_id))
        if section is None:
            raise HTTPException(status_code=404, detail="Section not found in this school")
        db.add(Enrollment(institution_id=user.institution_id, student_id=new_user.id, section_id=section.id, academic_year=section.academic_year))
    db.commit()
    return {"user": user_summary(new_user), "temporary_pin": temp_pin, "notice": "Share once; user must change it at first login."}


@app.post("/admin/users/import-csv", status_code=201)
def import_users(user: Annotated[User, Depends(require_roles(Role.ADMIN))], db: Db, file: UploadFile = File(...)) -> dict:
    try:
        filename = (file.filename or "").lower()
        raw = file.file.read()
        if filename.endswith(".csv"):
            reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
            fieldnames = [field.strip().lower() for field in (reader.fieldnames or [])]
            rows = [dict(zip(fieldnames, row.values())) for row in reader]
        elif filename.endswith(".xlsx"):
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            worksheet = workbook.active
            all_rows = list(worksheet.iter_rows(values_only=True))
            fieldnames = [str(value).strip().lower() for value in (all_rows[0] if all_rows else [])]
            rows = [dict(zip(fieldnames, values)) for values in all_rows[1:]]
        else:
            raise HTTPException(status_code=400, detail="Upload a .csv or .xlsx roster")
        required = {"full_name", "email", "role"}
        if not fieldnames or not required.issubset(set(fieldnames)):
            raise HTTPException(status_code=400, detail="Roster must include full_name, email, and role columns")
        created = []
        for row in rows:
            if not row.get("full_name") or not row.get("email") or not row.get("role"):
                continue
            role = Role(str(row["role"]).strip().upper())
            if role == Role.COMMUNITY:
                raise ValueError("Community accounts cannot be bulk-created")
            email = str(row["email"]).strip().lower()
            if db.scalar(select(User.id).where(User.institution_id == user.institution_id, User.email == email)):
                continue
            pin = secrets.token_urlsafe(9)
            item = User(institution_id=user.institution_id, full_name=str(row["full_name"]).strip(), email=email,
                        role=role.value, password_hash=hash_password(pin), force_password_change=True)
            db.add(item)
            db.flush()
            created.append({"email": email, "temporary_pin": pin, "role": role.value})
        db.commit()
        return {"created": len(created), "credentials": created, "note": "Temporary credentials are returned once; securely deliver them to users."}
    except HTTPException:
        raise
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"CSV import failed: {exc}")


@app.get("/admin/users")
def list_users(user: Annotated[User, Depends(require_roles(Role.ADMIN))], db: Db, role: Role | None = None) -> list[dict]:
    query = select(User).where(User.institution_id == user.institution_id)
    if role:
        query = query.where(User.role == role.value)
    return [user_summary(item) for item in db.scalars(query.order_by(User.full_name)).all()]


@app.post("/admin/sections", status_code=201)
def create_section(payload: SectionCreate, user: Annotated[User, Depends(require_roles(Role.ADMIN))], db: Db) -> dict:
    if payload.head_teacher_id:
        head = db.scalar(select(User).where(User.id == payload.head_teacher_id, User.institution_id == user.institution_id, User.role == Role.TEACHER.value))
        if head is None:
            raise HTTPException(status_code=400, detail="Head teacher must be a teacher in this school")
    section = Section(institution_id=user.institution_id, **payload.model_dump())
    db.add(section)
    db.commit()
    return {"id": section.id, "grade_level": section.grade_level, "section_name": section.section_name,
            "academic_year": section.academic_year, "head_teacher_id": section.head_teacher_id}


@app.get("/sections")
def list_sections(user: CurrentUser, db: Db) -> list[dict]:
    if user.role == Role.STUDENT.value:
        query = select(Section).join(Enrollment, Enrollment.section_id == Section.id).where(Enrollment.student_id == user.id)
    else:
        query = select(Section).where(Section.institution_id == user.institution_id)
    return [{"id": s.id, "grade_level": s.grade_level, "section_name": s.section_name,
             "academic_year": s.academic_year, "head_teacher_id": s.head_teacher_id} for s in db.scalars(query).all()]


@app.post("/marklists", status_code=201)
def create_marklist(payload: MarklistCreate, user: Annotated[User, Depends(require_roles(Role.ADMIN, Role.TEACHER))], db: Db) -> dict:
    section = db.scalar(select(Section).where(Section.id == payload.section_id, Section.institution_id == user.institution_id))
    subject = db.scalar(select(Subject).where(Subject.id == payload.subject_id, Subject.institution_id == user.institution_id))
    if not section or not subject:
        raise HTTPException(status_code=404, detail="Section or subject not found in this school")
    if user.role == Role.TEACHER.value:
        assigned = db.scalar(select(TeachingAssignment.id).where(TeachingAssignment.teacher_id == user.id,
                      TeachingAssignment.section_id == section.id, TeachingAssignment.subject_id == subject.id))
        if not assigned:
            raise HTTPException(status_code=403, detail="Teachers may only create marklists for assigned subjects and sections")
    item = Marklist(institution_id=user.institution_id, teacher_id=user.id, **payload.model_dump())
    db.add(item)
    db.commit()
    return {"id": item.id, "title": item.title, "max_points": item.max_points, "published": item.published}


@app.put("/marklists/{marklist_id}/scores")
def update_scores(marklist_id: str, payload: ScoreInput, user: Annotated[User, Depends(require_roles(Role.ADMIN, Role.TEACHER))], db: Db) -> dict:
    marklist = db.scalar(select(Marklist).where(Marklist.id == marklist_id, Marklist.institution_id == user.institution_id))
    if marklist is None:
        raise HTTPException(status_code=404, detail="Marklist not found")
    if user.role == Role.TEACHER.value:
        assigned = db.scalar(select(TeachingAssignment.id).where(TeachingAssignment.teacher_id == user.id,
                      TeachingAssignment.section_id == marklist.section_id, TeachingAssignment.subject_id == marklist.subject_id))
        if marklist.teacher_id != user.id and not assigned:
            raise HTTPException(status_code=403, detail="You are not assigned to this marklist")
    student_ids = set(db.scalars(select(Enrollment.student_id).where(Enrollment.institution_id == user.institution_id,
                         Enrollment.section_id == marklist.section_id, Enrollment.academic_year == marklist.academic_year)).all())
    for student_id, points in payload.scores.items():
        if student_id not in student_ids:
            raise HTTPException(status_code=400, detail=f"Student {student_id} is not enrolled in this section and year")
        if points < 0 or points > marklist.max_points:
            raise HTTPException(status_code=400, detail=f"Score must be between 0 and {marklist.max_points}")
        score = db.scalar(select(Score).where(Score.marklist_id == marklist.id, Score.student_id == student_id))
        if score:
            score.earned_points = points
        else:
            db.add(Score(institution_id=user.institution_id, marklist_id=marklist.id, student_id=student_id, earned_points=points))
    marklist.published = payload.publish
    db.commit()
    return {"message": "Scores saved", "published": marklist.published, "updated": len(payload.scores),
            "rating_rule": "sum(earned_points) / sum(max_points) * 10"}


@app.get("/performance/me")
def my_performance(user: Annotated[User, Depends(require_roles(Role.STUDENT))], db: Db) -> dict:
    section = student_section(db, user)
    if section is None:
        return {"earned_points": 0, "maximum_points": 0, "average_percent": None, "calculated_stars": None, "scale": 10.0}
    rows = db.execute(select(func.coalesce(func.sum(Score.earned_points), 0.0),
                              func.coalesce(func.sum(Marklist.max_points), 0.0))
                      .select_from(Marklist)
                      .outerjoin(Score, (Score.marklist_id == Marklist.id) & (Score.student_id == user.id))
                      .where(Marklist.institution_id == user.institution_id,
                             Marklist.section_id == section.id,
                             Marklist.academic_year == section.academic_year,
                             Marklist.published.is_(True))).one()
    earned, maximum = float(rows[0]), float(rows[1])
    stars = round((earned / maximum) * 10, 2) if maximum > 0 else None
    return {"earned_points": earned, "maximum_points": maximum, "average_percent": round((earned / maximum) * 100, 2) if maximum else None,
            "calculated_stars": stars, "scale": 10.0}


@app.post("/posts", status_code=201)
def create_post(payload: PostCreate, user: CurrentUser, db: Db) -> dict:
    if user.role not in {Role.ADMIN.value, Role.TEACHER.value}:
        raise HTTPException(status_code=403, detail="Only administrators and teachers can publish scoped posts")
    if payload.scope == Scope.GLOBAL and user.role != Role.ADMIN.value:
        raise HTTPException(status_code=403, detail="Only administrators can publish global announcements")
    if payload.scope == Scope.GRADE and payload.grade_level is None:
        raise HTTPException(status_code=400, detail="grade_level is required for a grade feed")
    if payload.scope == Scope.SECTION and not payload.section_id:
        raise HTTPException(status_code=400, detail="section_id is required for a section feed")
    if payload.section_id:
        section = db.scalar(select(Section).where(Section.id == payload.section_id, Section.institution_id == user.institution_id))
        if section is None:
            raise HTTPException(status_code=404, detail="Section not found")
        require_section_access(db, user, section.id)
    post = Post(institution_id=user.institution_id, sender_id=user.id, scope=payload.scope.value,
                grade_level=payload.grade_level, section_id=payload.section_id,
                title=payload.title, content=payload.content)
    db.add(post)
    db.commit()
    return {"id": post.id, "scope": post.scope, "title": post.title, "content": post.content}


@app.get("/posts")
def list_posts(user: CurrentUser, db: Db) -> list[dict]:
    query = select(Post).where(Post.institution_id == user.institution_id)
    if user.role == Role.STUDENT.value:
        section = student_section(db, user)
        conditions = [Post.scope == Scope.GLOBAL.value]
        if section:
            conditions += [Post.scope == Scope.GRADE.value, Post.scope == Scope.SECTION.value]
            query = query.where((Post.scope == Scope.GLOBAL.value) | ((Post.scope == Scope.GRADE.value) & (Post.grade_level == section.grade_level)) |
                                ((Post.scope == Scope.SECTION.value) & (Post.section_id == section.id)))
        else:
            query = query.where(Post.scope == Scope.GLOBAL.value)
    elif user.role == Role.COMMUNITY.value:
        query = query.where(Post.scope == Scope.GLOBAL.value)
    else:
        query = query
    rows = db.scalars(query.order_by(Post.created_at.desc()).limit(100)).all()
    return [{"id": p.id, "scope": p.scope, "title": p.title, "content": p.content,
             "created_at": p.created_at.isoformat()} for p in rows]


@app.post("/messages/difficulty", status_code=201)
def report_difficulty(payload: MessageInput, user: Annotated[User, Depends(require_roles(Role.STUDENT))], db: Db) -> dict:
    section = student_section(db, user)
    if section is None or section.head_teacher_id is None:
        raise HTTPException(status_code=409, detail="No Class Head Teacher is assigned to your section")
    message = Post(institution_id=user.institution_id, sender_id=user.id, recipient_id=section.head_teacher_id,
                   scope=Scope.DM.value, category="STUDENT_DIFFICULTY", title="Student support request", content=payload.content)
    db.add(message)
    db.commit()
    return {"message": "Your report was sent privately to your Class Head Teacher."}


@app.post("/messages/guidance", status_code=201)
def contact_guidance(payload: MessageInput, user: Annotated[User, Depends(require_roles(Role.STUDENT))], db: Db) -> dict:
    school = db.get(Institution, user.institution_id)
    if not school or not school.has_guidance_counseling:
        raise HTTPException(status_code=403, detail="Guidance & Counseling is not enabled for this school")
    counselor = db.scalar(select(User).where(User.institution_id == user.institution_id, User.role == Role.COUNSELOR.value, User.active.is_(True)))
    if counselor is None:
        raise HTTPException(status_code=409, detail="No Guidance & Counseling contact is currently configured")
    db.add(Post(institution_id=user.institution_id, sender_id=user.id, recipient_id=counselor.id,
                scope=Scope.DM.value, category="GUIDANCE_SUPPORT", title="Student support", content=payload.content))
    db.commit()
    return {"message": "Your private message was sent to Guidance & Counseling."}


@app.post("/messages/community-inquiry", status_code=201)
def community_inquiry(payload: MessageInput, user: Annotated[User, Depends(require_roles(Role.COMMUNITY))], db: Db) -> dict:
    admin = db.scalar(select(User).where(User.institution_id == user.institution_id, User.role == Role.ADMIN.value, User.active.is_(True)))
    if admin is None:
        raise HTTPException(status_code=409, detail="No active school administrator is available")
    db.add(Post(institution_id=user.institution_id, sender_id=user.id, recipient_id=admin.id,
                scope=Scope.DM.value, category="PEOPLE_INQUIRY", title="Community inquiry", content=payload.content))
    db.commit()
    return {"message": "Your inquiry was sent to the school administrator."}


@app.get("/admin/inbox")
def admin_inbox(user: Annotated[User, Depends(require_roles(Role.ADMIN))], db: Db, category: str | None = None) -> list[dict]:
    query = select(Post).where(Post.institution_id == user.institution_id, Post.scope == Scope.DM.value)
    if category:
        query = query.where(Post.category == category.upper())
    rows = db.scalars(query.order_by(Post.created_at.desc()).limit(200)).all()
    return [{"id": p.id, "sender_id": p.sender_id, "recipient_id": p.recipient_id, "category": p.category,
             "title": p.title, "content": p.content, "created_at": p.created_at.isoformat()} for p in rows]


@app.get("/exports/marklists/{marklist_id}.xlsx")
def export_marklist(marklist_id: str, user: CurrentUser, db: Db):
    marklist = db.scalar(select(Marklist).where(Marklist.id == marklist_id, Marklist.institution_id == user.institution_id))
    if marklist is None:
        raise HTTPException(status_code=404, detail="Marklist not found")
    require_section_access(db, user, marklist.section_id)
    if user.role == Role.STUDENT.value and not marklist.published:
        raise HTTPException(status_code=404, detail="Marklist not found")
    score_query = select(User.full_name, User.email, Score.earned_points).join(Score, Score.student_id == User.id).where(
        Score.marklist_id == marklist.id, Score.institution_id == user.institution_id)
    if user.role == Role.STUDENT.value:
        score_query = score_query.where(Score.student_id == user.id)
    scores = db.execute(score_query.order_by(User.full_name)).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Marklist"
    sheet.append([marklist.title, "Maximum points", marklist.max_points])
    sheet.append(["Student", "Email", "Earned points", "Maximum points", "Percentage"])
    for name, email, points in scores:
        sheet.append([name, email, points, marklist.max_points, round(points / marklist.max_points * 100, 2)])
    output = io.BytesIO()
    workbook.save(output)
    from fastapi.responses import StreamingResponse
    output.seek(0)
    filename = f"schoolnet-{marklist.academic_year}-{marklist.title[:40].replace(' ', '-')}.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


# Register additional administrator, inbox and report-card routes after core services load.
from . import extensions as _extensions  # noqa: E402,F401
